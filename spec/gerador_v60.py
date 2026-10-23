# -*- coding: utf-8 -*-
"""
gerador_v60.py — regera a cadeia v60 (23/out/2026) a partir dos arquivos v59 do projeto + snapshot da aba Tiers.

Entradas (na pasta corrente): playbook_v59.md · fichas_tecnicas_v9.md · tabelao_v2_7.md/.xlsx · frascos_v7.py ·
modelo_diario_v1_6.py · observatorio_colecao_v4_6.html · mapa_climatico_colecao_v5.html ·
avaliacao_perfumes_v51_2026-10-05.html · 00_LEIA_PRIMEIRO.md.   Saídas: pasta out/.
Snapshot da aba Tiers do Google Sheets (23/out/2026, após a limpeza feita com ele) embutido em TIERS_TSV.

O que muda (e só isso — verificado por assert):
  · entra Blue Sapphire Supercharged (grade, S) e Empire (camada custo, A)
  · Bedouin Rose D→C entra na grade (decisão dele): só células-lar
  · Favonius, Omnia Omnibus Ubique, Scarlet Sands: envelope/aptidão revistos pelo Claude a pedido dele (estimativa)
  · Seguros, jogadas e sprays de Seguro idênticos à v59 (incumbência); nenhum outro frasco muda de célula
  · slot E (Esporte) e NotaDia (acerto do dia, 1–5) formalizados no LEIA_PRIMEIRO; regra do doce (18/out) incorporada
Depois de rodar: python3 recalc.py out/tabelao_v2_8.xlsx (recalcula as fórmulas da aba Distribuição).
"""
TIERS_TSV = """Abdul Samad Al Qurashi|Al Qurashi Blend|Al Qurashi Blend|Frasco Gui|A||
Aesop|Above Us, Steorra||Avaliado|B|2026-10-12|
Afnan|Supremacy CE|Supremacy CE|Frasco Gui|S||
Al Ambra|Dubai Musk|Dubai Musk|Frasco Gui|A||
Amouage|Beach Hut Man|Beach Hut Man|Frasco Gui|A||
Amouage|Decision|Decision|Frasco Gui|S||
Amouage|Elsewhere||Avaliado|C|2026-10-23|
Amouage|Epic Man||Amostra|B|2026-10-12|
Amouage|Guidance 46||Frasco Bia|S||
Amouage|Honour Woman||Amostra Bia|A||
Amouage|Interlude 53||Amostra|B||
Amouage|Interlude Man||Avaliado|B||
Amouage|Jubilation 40||Amostra|S|2026-10-04|
Amouage|Love Hibiscus||Amostra Bia|B||
Amouage|Opus XIV Royal Tobacco (Library)||Avaliado|B||
Amouage|Opus XVI Timber (Library)||Avaliado|B||
Amouage|Outlands|Outlands|Frasco Gui|S||
Amouage|Purpose||Amostra|C||
Amouage|Purpose 50|Purpose 50|Frasco Gui|S||
Amouage|Reasons (Essences)||Avaliado|B||
Amouage|Reflection 45||Amostra|A||
Amouage|Reflection Man|Reflection Man|Frasco Gui|S||
Amouage|Remain (Essences)||Avaliado|B||
Amouage|Search|Search|Frasco Gui|A||
Amouage|Sequence||Amostra|A||
Andy Tauer|Au Coeur du Désert||Amostra|B||
Arabiyat Prestige|Marwa||Avaliado|B||
Armani Privé|Orange Méditerranée||Avaliado|B|2026-10-23|
Atelier des Ors|Blanc Polychrome||Avaliado|A|2026-10-23|
Atelier Versace|Gingembre Pétillant||Amostra|B||
Atelier Versace|Patchouli Précieux||Amostra|B||
Atelier Versace|Santal Boisé||Amostra|B||
BDK|312 Saint-Honoré||Avaliado|B||
BDK|Crème de Cuir||Avaliado|B||
BDK|Gris Charnel EDP||Avaliado|B||
BDK|Gris Charnel Extrait|Gris Charnel|Frasco Gui|S||
BDK|Rouge Smoking||Avaliado|B||
BDK|Rouge Smoking Extrait||Avaliado|B||
BDK|Silver Ceremony (2026)||Avaliado|A||
BDK|Tabac Rose||Avaliado|B||
Boadicea|Ardent||Amostra|B||
Boadicea|Blue Sapphire||Avaliado|A||
Boadicea|Bodacious||Amostra|B||
Boadicea|Dragon||Amostra|A||
Boadicea|Harmonious||Amostra|C||
Boadicea|Imperial|Imperial|Frasco Gui|A||
Boadicea|Knight of Love||Amostra|A||
Boadicea|Monarch|Monarch|Frasco Gui|B||
Boadicea|Seductive||Amostra|C||
Boadicea|Valiant||Amostra|A||
Born to Stand Out|Black Crème||Avaliado|A||
Born to Stand Out|Dirty Rice||Avaliado|C||
Born to Stand Out|Drunk Lovers|Drunk Lovers|Frasco Gui|A||
Born to Stand Out|Purple Stain||Avaliado|C||
Born to Stand Out|Sugar Addict||Avaliado|C||
Bvlgari|Le Gemme Tygar||Amostra|S||
Bvlgari|Le Gemme Tygar Extrait|Tygar Extrait|Frasco Gui|S||
By Kilian|Angels' Share||Amostra|B||
By Kilian|Intoxicated|Intoxicated|Frasco Gui|A||
BYNOMADS|Bedouin Rose|Bedouin Rose|Frasco Gui|C|2026-10-14|
BYNOMADS|K-Musk|K-Musk|Frasco Gui|B||
BYNOMADS|Oud Kyphi|Oud Kyphi|Frasco Gui|C||
Byredo|Alto Astral||Avaliado|B||
Byredo|Bal d'Afrique||Avaliado|B||
Byredo|Bal d'Afrique Absolu||Frasco Bia|S||
Byredo|Bibliothèque||Frasco Bia|A||
Byredo|Black Saffron||Avaliado|B||
Byredo|Black Saffron Absolu|Black Saffron Absolu|Frasco Gui|A||
Byredo|Blanche||Frasco Bia|B||
Byredo|Blanche Absolu||Avaliado|B||
Byredo|De Los Santos||Avaliado|C||
Byredo|Future Memories||Avaliado|B|2026-10-23|
Byredo|Gypsy Water||Avaliado|A|2026-10-23|
Byredo|Mojave Ghost Absolu||Avaliado|A||
Byredo|Oud Immortel||Avaliado|B||
Byredo|Rose of No Man's Land Absolu||Avaliado|A||
Byredo|Rouge Chaotique (Night Veils)||Avaliado|A||
Byredo|Vanille Antique||Avaliado|A||
Byredo|Vanille Antique 2025||Avaliado|A||
Carolina Herrera|Cedar Chic||Avaliado|B||
Chanel|Bleu de Chanel EDP|Bleu de Chanel|Frasco Gui|C||
Clive Christian|1872 for Men|1872 for Men|Frasco Gui|S||
Clive Christian|Blonde Amber (XXI Art Deco)|Blonde Amber|Frasco Gui|S||
Clive Christian|Crab Apple Blossom (Crown)|Crab Apple Blossom|Frasco Gui|S||
Clive Christian|E Cashmere Musk (Private)||Avaliado|B||
Clive Christian|Jump Up and Kiss Me Hedonistic||Amostra|B||
Clive Christian|Matsukita (Crown)||Avaliado|B||
Clive Christian|Strange Heavens Out of the Blue||Avaliado|B||
Clive Christian|Timeless (150 Anniversary)||Avaliado|A||
Clive Christian|Town & Country (Crown)|Town & Country|Frasco Gui|A||
Creed|Aventus|Aventus|Frasco Gui|A||
Creed|Aventus Absolu|Aventus Absolu|Frasco Gui|S||
Creed|Aventus Cologne||Amostra|S|2026-10-23|
Creed|Bois du Portugal|Bois du Portugal|Frasco Gui|B||
Creed|Centaurus||Avaliado|B||
Creed|Erolfa||Avaliado|B||
Creed|Green Irish Tweed|Green Irish Tweed|Frasco Gui|B||
Creed|Himalaya|Himalaya|Frasco Gui|B||
Creed|Iris Debonair||Avaliado|A|2026-10-23|
Creed|Original Santal||Avaliado|B||
Creed|Original Vetiver||Avaliado|A||
Creed|Oud Zarian||Avaliado|B||
Creed|Royal Oud|Royal Oud|Frasco Gui|A||
Creed|Royal Water||Avaliado|B||
Creed|Viking||Avaliado|C||
Creed|Viking Cologne||Avaliado|C||
Creed|Wild Vetiver|Wild Vetiver|Frasco Gui|A||
D.S. & Durga|Debaser||Avaliado|B|2026-10-23|
Dior|Ambre Nuit||Avaliado|B||
Dior|Bois d'Argent||Avaliado|B||
Dior|Bois Talisman||Avaliado|B||
Dior|Cuir Saddle||Avaliado|B||
Dior|Dior Paradise||Avaliado|A||
Dior|Gris Dior||Amostra|B||
Dior|Vanilla Diorama||Avaliado|B||
Diptyque|L'Eau Papier||Avaliado|C||
Diptyque|Orphéon EDP||Avaliado|B||
Diptyque|Philosykos EDT||Avaliado|C||
Diptyque|Tam Dao EDP||Avaliado|B||
Diptyque|Vetyverio EDP||Avaliado|B||
Eight & Bob|Eight & Bob (Original)||Avaliado|B|2026-10-23|
Electimuss|Amber Bloom||Avaliado|A||
Electimuss|Astrum Nova||Avaliado|B||
Essential Parfums|Bois Impérial|Bois Impérial|Frasco Gui|S||
Ex Nihilo|Amber Sky||Avaliado|B||
Ex Nihilo|Atlas Fever||Avaliado|B||
Ex Nihilo|Blue Talisman EDP|Blue Talisman|Frasco Gui|A||
Ex Nihilo|Blue Talisman Extrait||Amostra|S||
Ex Nihilo|Bois d'Hiver||Avaliado|B||
Ex Nihilo|Brompton Immortals EDP||Avaliado|A||
Ex Nihilo|Brompton Immortals Extrait|Brompton|Frasco Gui|S||
Ex Nihilo|Chandigarh Express||Avaliado|A||
Ex Nihilo|Fleur Narcotique||Amostra|B||
Ex Nihilo|Fleur Narcotique Extrait||Avaliado|A||
Ex Nihilo|Lust in Paradise Extrait||Amostra|A||
Ex Nihilo|Outcast Blue||Avaliado|A||
Ex Nihilo|Rose Hubris||Avaliado|B||
Ex Nihilo|Santal Calling||Avaliado|B||
Ex Nihilo|Scarlet Sands|Scarlet Sands|Frasco Gui|S||
Ex Nihilo|Speed Legends Living on the Edge||Amostra|A|2026-10-23|
Ex Nihilo|The Hedonist||Avaliado|B||
Ex Nihilo|The Hedonist Extrait||Avaliado|A||
Fragrance du Bois|London Spice||Avaliado|A||
Fragrance du Bois|Lucius (regular)|Lucius|Frasco Gui|S||
Fragrance du Bois|Sahraa||Avaliado|B||
Frédéric Malle|Bigarade Concentrée||Avaliado|B||
Frédéric Malle|Carnal Flower||Frasco Bia|C||
Frédéric Malle|Cologne Indélébile||Avaliado|B||
Frédéric Malle|Dawn||Avaliado|B||
Frédéric Malle|French Lover||Avaliado|B||
Frédéric Malle|Hope||Avaliado|A||
Frédéric Malle|Musc Ravageur||Avaliado|A||
Frédéric Malle|Portrait of a Lady||Avaliado|B||
Frédéric Malle|Promise||Avaliado|C||
Frédéric Malle|The Moon||Avaliado|C||
Frédéric Malle|The Night||Avaliado|C||
Givenchy|Désobéissant|Désobéissant|Frasco Gui|S||
Granado|Apotecário (Vintage)|Apotecário|Frasco Gui|B||
Granado|Bossa||Frasco Bia|B||
Guerlain|Herbes Troublantes||Amostra|S|2026-10-23|
Guerlain|Oud Khôl||Avaliado|C|2026-10-23|
Guerlain|Oud Nude||Avaliado|C|2026-10-23|
Guerlain|Tobacco Honey||Avaliado|C|2026-10-23|
Hermès|Terre d'Hermès Parfum|Terre d'Hermès|Frasco Gui|A||
Hind Al Oud|Ya'e|Ya'e|Frasco Gui|B||
Initio|Absolute Aphrodisiac||Avaliado|B||
Initio|Atomic Rose||Amostra|B||
Initio|Blessed Baraka||Avaliado|B||
Initio|Can't Get Enough (2025)||Avaliado|B||
Initio|Musk Therapy||Frasco Bia|A||
Initio|Narcotic Delight||Avaliado|C||
Initio|Oud for Greatness||Amostra|B||
Initio|Paragon||Avaliado|C||
Initio|Rehab||Amostra|A||
Initio|Side Effect||Avaliado|B||
Initio|Sugar Blast (2026)||Avaliado|D||
Initio|Wild Rush||Amostra|B||
Jovoy|21 Conduit St||Avaliado|S|2026-10-23|
Khadlaj|Island Dreams|Island Dreams|Frasco Gui|B||
Khadlaj|Onyx Gold||Avaliado|B||
Khadlaj|Sky|Sky|Frasco Gui|A||
Lattafa|Breeze||Avaliado|B||
Lattafa|Dynasty||Avaliado|A||
Lattafa|Musamam Black|Musamam Black|Frasco Gui|A||
Le Labo|Ambrette 9||Frasco Bia|D||
Le Labo|Another 13||Avaliado|A||
Le Labo|Bergamote 22||Frasco Bia|B||
Le Labo|Santal 33|Santal 33|Frasco Gui|A||
Le Labo|Thé Noir 29||Avaliado|B||
Les Indémodables|Vanille Havane||Avaliado|B|2026-10-23|
Loewe|7 Loewe|7 Loewe|Frasco Gui|B||
Loewe|Solo||Amostra|B||
Louis Vuitton|Afternoon Swim||Frasco Bia|A||
Louis Vuitton|Ambre Levant (2026)||Avaliado|C||
Louis Vuitton|Attrape-Rêves||Frasco Bia|A||
Louis Vuitton|California Dream||Avaliado|B||
Louis Vuitton|City of Stars||Avaliado|B||
Louis Vuitton|Fleur du Désert||Avaliado|A||
Louis Vuitton|Imagination|Imagination|Frasco Gui|S||
Louis Vuitton|L'Immensité|L'Immensité|Frasco Gui|A||
Louis Vuitton|Les Sables Roses|Les Sables Roses|Frasco Gui|S||
Louis Vuitton|LVERS (2025, Pharrell)||Avaliado|B||
Louis Vuitton|Météore|Météore|Frasco Gui|S||
Louis Vuitton|Ombre Nomade|Ombre Nomade|Frasco Gui|S||
Louis Vuitton|Orage||Avaliado|B||
Louis Vuitton|Pacific Chill||Avaliado|B||
Louis Vuitton|Sun Song||Avaliado|B||
Louis Vuitton|Sur la Route||Avaliado|B||
Louis Vuitton|Symphony (Extraits)||Frasco Bia|A||
Maison Alhambra|Jean Lowe Vibe||Avaliado|C||
Maison Asrar|Belgravia|Belgravia|Frasco Gui|A||
Maison Crivelli|Hibiscus Mahajád||Avaliado|A||
Maison Crivelli|Iris Malikhân||Avaliado|B||
Maison Crivelli|Musc Nurāsana||Avaliado|A||
Maison Crivelli|Oud Maracuja||Amostra|A||
Maison Crivelli|Tobacco Carnaval||Avaliado|C||
Maison Margiela|By the Fireplace||Amostra|B||
Maison Margiela|Jazz Club||Amostra|A||
Maison Margiela|Lazy Sunday Morning||Avaliado|A||
Maison Margiela|Sailing Day||Avaliado|B||
Marrakech Imperial|Bohème Chic||Avaliado|B|2026-10-23|
Matière Première|Cologne Cédrat||Avaliado|B||
Matière Première|Crystal Saffron EDP||Amostra|A|2026-10-17|
Matière Première|Crystal Saffron Extrait||Avaliado|B||
Matière Première|Encens Suave||Avaliado|B||
Matière Première|Encens Suave Extrait||Avaliado|B||
Matière Première|Falcon Leather||Avaliado|B||
Matière Première|Falcon Leather Extrait||Avaliado|B||
Matière Première|Metal Lavender||Avaliado|C|2026-10-17|
Matière Première|Neroli Oranger||Amostra Bia|D||
Matière Première|Parisian Musc||Avaliado|B||
Matière Première|Radical Rose EDP||Amostra Bia|A||
Matière Première|Radical Rose Extrait|Radical Rose|Frasco Gui|S||
Matière Première|Santal Austral||Avaliado|B||
Matière Première|Santal Austral Extrait||Avaliado|B||
Matière Première|Vanilla Powder||Frasco Bia|D||
Matière Première|Vanilla Powder Extrait||Avaliado|D||
Memo Paris|African Leather|African Leather|Frasco Gui|S||
Memo Paris|African Rose||Avaliado|A||
Memo Paris|Cap Camarat||Avaliado|B||
Memo Paris|Ithaque||Avaliado|B||
Memo Paris|Madurai||Amostra Bia|B||
Memo Paris|Marfa||Amostra|B||
Memo Paris|Ocean Leather (Cuirs Nomades)||Amostra|B||
Memo Paris|Odéon|Odéon|Frasco Gui|S||
Memo Paris|Odéon Rosebud (2026)||Avaliado|C||
Memo Paris|Omnia Omnibus Ubique|Omnia Omnibus Ubique|Frasco Gui|S||
MFK|724||Avaliado|B||
MFK|Amyris Homme Extrait|Amyris Homme Extrait|Frasco Gui|S||
MFK|APOM||Avaliado|B||
MFK|Gentle Fluidity Silver|Gentle Fluidity Silver|Frasco Gui|S||
MFK|Grand Soir|Grand Soir|Frasco Gui|S||
MFK|Masculin Pluriel||Amostra|A|2026-10-23|
MFK|Oud (extrait)||Avaliado|B||
MFK|Oud Satin Mood Extrait|Oud Satin Mood|Frasco Gui|S||
MFK|Oud Silk Mood||Avaliado|B||
MFK|Oud Velvet Mood||Avaliado|D||
Mind Games|Queening||Avaliado|D||
Mykonos|Inception|Inception|Frasco Gui|S||
Nasomatto|Black Afgano||Avaliado|D|2026-10-23|
Nishane|Ani||Avaliado|B||
Nishane|Ani X|Ani X|Frasco Gui|A||
Nishane|Favonius|Favonius|Frasco Gui|S||
Nishane|Hacivat|Hacivat|Frasco Gui|S||
Nishane|Hundred Silent Ways||Avaliado|C||
Nishane|Kredo||Avaliado|B||
Nishane|Meant To Be Seen (2025)||Avaliado|C||
Nishane|Oudous Lux Solis|Oudós Lux Solis|Frasco Gui|A||
Nishane|Papilefiko||Avaliado|B||
Nishane|Tero||Avaliado|B||
Nishane|Thorns Become Roses||Avaliado|B||
Nishane|Wulong Cha||Avaliado|A||
OJAR|Wood Whisper|Wood Whisper|Frasco Gui|A||
Omanluxury|Caden|Caden|Frasco Gui|A||
Omanluxury|Nasaj||Amostra|B|2026-10-16|
Omanluxury|Oud Aquilaria||Avaliado|B||
Omanluxury|Overdose||Avaliado|B||
Omanluxury|Paramour||Avaliado|B||
Omanluxury|Royal Incense||Amostra|C|2026-10-16|
Omanluxury|Serenity|Serenity|Frasco Gui|B||
Omanluxury|Wanderlust 2023||Avaliado|B||
Ormonde Jayne|Bijou Zafran|Bijou Zafran|Frasco Gui|S||
Ormonde Jayne|Montabaco Intensivo|Montabaco Intensivo|Frasco Gui|A||
Orto Parisi|Bergamask||Avaliado|C||
Orto Parisi|Megamare||Avaliado|C||
Orto Parisi|Risvelium||Avaliado|D||
Pantheon Roma|Annone||Avaliado|B||
Parfums de Marly|Castley|Castley|Frasco Gui|B||
Parfums de Marly|Delina||Frasco Bia|B||
Parfums de Marly|Galloway||Avaliado|B||
Parfums de Marly|Greenley||Amostra|B||
Parfums de Marly|Haltane||Amostra|A||
Parfums de Marly|Layton||Amostra|A||
Parfums de Marly|Percival||Avaliado|B||
Parfums de Marly|Sedley||Avaliado|B||
Parfums de Marly|Valaya||Frasco Bia|A||
Penhaligon's|Cairo||Avaliado|S||
Penhaligon's|Empressa||Avaliado|D||
Penhaligon's|Halfeti||Avaliado|S||
Penhaligon's|Halfeti Leather||Avaliado|A||
Penhaligon's|The Blazing Mister Sam||Avaliado|A||
Pernoire|Anelo (Odus)|Anelo|Frasco Gui|A||
Pernoire|Yuma||Avaliado|B||
Profumum Roma|Acqua Viva||Amostra|A||
Rayhaan|Aquatica|Aquatica|Frasco Gui|B||
Regalien|Turkuaz||Amostra|B||
Riiffs|Freeze||Avaliado|B||
Roja Parfums|Amber Aoud||Avaliado|S||
Roja Parfums|Aoud||Avaliado|B||
Roja Parfums|Apex||Avaliado|A||
Roja Parfums|Burlington 1819||Avaliado|S||
Roja Parfums|Elysium Pour Homme||Avaliado|A||
Sospiro|Vibrato|Vibrato|Frasco Gui|A||
Swiss Arabian|Incense 01|Incense 01|Frasco Gui|B||
Thameen|Carved Oud|Carved Oud|Frasco Gui|A||
Thameen|Regent Leather||Amostra|B||
Thomas de Monaco|Ultima Storia|Ultima Storia|Frasco Gui|A||
Tom Ford|Tuscan Leather||Avaliado|A||
Widian|Liwa|Liwa|Frasco Gui|S||
Widian|London||Avaliado|A||
Widian|New York|New York|Frasco Gui|A||
Widian|Yasat||Avaliado|A||
Xerjoff|40 Knots||Avaliado|A||
Xerjoff|Alexandria II||Avaliado|B||
Xerjoff|Erba Gold||Avaliado|D||
Xerjoff|Erba Pura||Avaliado|D||
Xerjoff|More Than Words||Avaliado|B||
Xerjoff|Nio||Avaliado|A||
Xerjoff|Richwood||Avaliado|B||
Xerjoff|Torino 21|Torino 21|Frasco Gui|A||
Xerjoff|XJ 1861 Naxos||Avaliado|B||
Xerjoff|XJ 1861 Renaissance||Avaliado|B||
Yves Saint Laurent|Tuxedo||Avaliado|A||
Fragrance du Bois|Oud Bleu Intense||Avaliado|B|2026-10-23|
Memo Paris|Portobello Road||Amostra|C|2026-10-23|
Boadicea|Blue Sapphire Supercharged|Blue Sapphire Supercharged|Frasco Gui|S|2026-10-09|Sem ficha técnica; fora do playbook e do motor até a próxima geração da grade
Mykonos|Empire|Empire|Frasco Gui|A|2026-10-11|Camada custo (só corrida/academia); sem ficha técnica; fora do motor até a próxima geração
Etat Libre d’Orange|You Or Someone Like||Amostra|B|2026-10-12|
Les EAUX Primordiales|Ambre Supermassive||Amostra|S|2026-10-22|
Amouage|Bracken Man||Avaliado|S|2026-10-23|
Amouage|Boundless||Amostra|C|2026-10-16|
Amouage|Meander||Amostra|B|2026-10-13|
Amouage|Enclave||Amostra||2026-10-12|Tier pendente (ainda vai avaliar)
Amouage|Lyric Man||Amostra|B|2026-10-15|
Amouage|Honour Man||Amostra|C|2026-10-15|
Fragrance du Bois|Cavort||Amostra Bia|B|2026-10-12|
Arabiyat Prestige|Mahd Al Dhahab||Amostra|A|2026-10-23|
Mykonos|Dreamscape||Frasco Bia|A|2026-10-23|
Amouage|Myths Man||Amostra|C|2026-10-23|
Amouage|Journey Man||Amostra|B|2026-10-23|
Boadicea|Lannosea||Amostra|B|2026-10-23|
Fragrance du Bois|Secret Tryst||Amostra|A|2026-10-23|
Tom Ford|Beau||Amostra|B|2026-10-17|
Sora Dora|Gladiator||Avaliado|B|2026-10-23|
Min New York|Voodoo||Avaliado|A|2026-10-23|
"""
import re
BANDS = [('🔴', 'Muito Quente'), ('🟠', 'Quente'), ('🟡', 'Ameno'), ('🔵', 'Fresco'), ('🟣', 'Frio'), ('⚫', 'Muito Frio')]
OCC = ['Lazer', 'Trab. Informal', 'Trab. Formal', 'Cozy Dia', 'Noite Informal', 'Noite Formal', 'Cozy Noite']
OCC_AB = {'Lazer': 'Lazer', 'Trab. Informal': 'T.Inf', 'Trab. Formal': 'T.For', 'Cozy Dia': 'Cozy D',
          'Noite Informal': 'N.Inf', 'Noite Formal': 'N.For', 'Cozy Noite': 'Cozy N'}
CELL_RE = re.compile(r'^\*\*(?P<occ>[^*]+)\*\* — \*\*Seguro: (?P<s>[^*]+)\*\* \((?P<sp>\d+)\)\. (?P<sdesc>.*?) — tier (?P<st>[SABCD])\. '
                     r'\*\*A jogada: (?P<j>[^*]+)\*\* — (?P<jr>.*?) \(tier (?P<jt>[SABCD])\)\.(?: \*\*Também:\*\* (?P<t>.*)\.)?$')

def parse(path):
    L = open(path, encoding='utf-8').read().split('\n')
    cells = {}; band = None; where = {}
    for i, l in enumerate(L):
        m = re.match(r'^## (\S+) (.+?) \(', l)
        if m and any(m.group(1) == e for e, _ in BANDS):
            band = BANDS.index((m.group(1), m.group(2))); continue
        if l.startswith('## ') : band = None if not m else band
        if band is not None and l.startswith('**') and ' — **Seguro:' in l:
            c = CELL_RE.match(l)
            assert c, l[:200]
            g = c.groupdict()
            tam = [x.strip() for x in g['t'].split(',')] if g['t'] else []
            cells[(band, g['occ'])] = dict(s=g['s'], sp=int(g['sp']), sdesc=g['sdesc'], st=g['st'],
                                         j=g['j'], jr=g['jr'], jt=g['jt'], t=tam)
            where[(band, g['occ'])] = i
    return L, cells, where


# ======== PARTE A — dados da decisão (23/out/2026) + playbook_v60.md ========
import re, collections, json, os
os.makedirs('out', exist_ok=True)
def rep1(s, a, b, n=1):
    assert s.count(a) == n, (s.count(a), a[:90]); return s.replace(a, b)

# ---- tiers: aba Tiers do Sheets, snapshot de 23/out/2026 (após a limpeza feita com ele)
TIERS = [l.split('|') for l in TIERS_TSV.strip('\n').split('\n')]
assert len(TIERS) == 355
DIARIO_TIER = {r[2]: r[4] for r in TIERS if r[3] == 'Frasco Gui'}
assert len(DIARIO_TIER) == 81

# ---- base atual (frascos_v7)
import importlib.util
spec = importlib.util.spec_from_file_location('fv7', 'frascos_v7.py'); fv7 = importlib.util.module_from_spec(spec); spec.loader.exec_module(fv7)
MAPA = {m['n']: dict(m) for m in fv7.MAPA}
FICHA = {k: list(v) for k, v in fv7.FICHA.items()}

# ---- decisões de 23/out/2026 (envelope e aptidão: Claude, pela pirâmide + Fragrantica/Parfumo/marca, a pedido dele;
#      tier e posse: dele; Bedouin Rose C na grade: dele; slot E e NotaDia: dele)
# envelope: a = faixa mais quente, b = faixa mais fria (0 MQ · 1 Q · 2 A · 3 F · 4 Fr · 5 MF)
NOVOS_GRADE = {
 'Blue Sapphire Supercharged': dict(casa='Boadicea', arq='Rosal', t='S', td='S', p=4, d=1, a=2, b=5, c=0, r=0.55, u=0,
     ficha=('Parfum', 'S', 'Lazer 1 · N.Inf 3 · N.For 3 · Cozy N 1', 3)),
}
NOVOS_CUSTO = {
 'Empire': dict(casa='Mykonos', arq='Marinho', t='A', td='C', p=2, d=1, a=0, b=2, c=1, r=-0.5, u=0,
     ficha=('n/d', 'F', 'Lazer 3 · T.Inf 1', 4)),
}
MUDA = {  # nome: (campos MAPA, aptidão)
 'Favonius': (dict(a=2, b=5), 'N.Inf 3 · N.For 3 · Cozy N 2'),
 'Omnia Omnibus Ubique': (dict(a=2, b=5), 'N.Inf 2 · N.For 3 · Cozy N 3'),
 'Scarlet Sands': (dict(a=1, b=4), 'Lazer 2 · N.Inf 3 · N.For 2'),
 'Bedouin Rose': (dict(a=4, b=5, t='C'), 'Cozy N 3'),
}
for n, (mf, apt) in MUDA.items():
    MAPA[n].update(mf); FICHA[n][2] = apt
for n, d in {**NOVOS_GRADE, **NOVOS_CUSTO}.items():
    f = d.pop('ficha'); MAPA[n] = dict(n=n, **d); FICHA[n] = [f[0], f[1], f[2], f[3], None]
# tiers conferidos com o Sheets
for n, m in MAPA.items():
    assert DIARIO_TIER[n] == m['t'], (n, DIARIO_TIER[n], m['t'])
assert set(MAPA) == set(DIARIO_TIER)

def apt_of(n): return fv7.parse_aptidao(FICHA[n][2])
TRANK = {'S': 0, 'A': 1, 'B': 2, 'C': 3, 'D': 4}
CAP = {'S': 12, 'A': 9, 'B': 6, 'C': 3, 'D': 0}
# frequência do slot por faixa (normais INMET, já no playbook): desempate de fit = célula mais usada no ano
DIAS = [27, 124, 154, 55, 5, 0]; NOITES = [1, 36, 154, 143, 30, 1]
DIA_OCC = {'Lazer', 'Trab. Informal', 'Trab. Formal', 'Cozy Dia'}

def eligible(n, cells):
    m = MAPA[n]; apt = apt_of(n); mx = max(apt.values())
    out = []
    for (b, o) in cells:
        if not (m['a'] <= b <= m['b']): continue
        ap = apt.get(OCC_AB[o], 0)
        if ap < 1: continue
        if m['t'] == 'C' and ap < mx: continue
        ideal = (m['a'] + m['b']) / 2 + 0.5 * (m['p'] - 3)
        ideal = max(m['a'], min(m['b'], ideal))
        fit = ap - 0.5 * abs(b - ideal)
        freq = DIAS[b] if o in DIA_OCC else NOITES[b]
        out.append(((-fit, -freq, b, OCC.index(o)), (b, o)))
    out.sort()
    return [c for _, c in out[:CAP[m['t']]]]

L, cells, where = parse('playbook_v59.md')
assert sum(2 + len(c['t']) for c in cells.values()) == 651
ANTES = {k: [c['s'], c['j']] + c['t'] for k, c in cells.items()}

REALOCA = ['Favonius', 'Omnia Omnibus Ubique', 'Scarlet Sands', 'Blue Sapphire Supercharged', 'Bedouin Rose']
for c in cells.values():
    for n in REALOCA:
        assert n not in (c['s'], c['j'])  # incumbência: nenhum deles é Seguro/jogada
        if n in c['t']: c['t'].remove(n)
NOVAS = {n: eligible(n, cells) for n in REALOCA}

def key(n, o): return (TRANK[MAPA[n]['t']], -apt_of(n).get(OCC_AB[o], 0))
for n, cl in NOVAS.items():
    for (b, o) in cl:
        t = cells[(b, o)]['t']; k = key(n, o)
        pos = 0
        for i, x in enumerate(t):
            if key(x, o) <= k: pos = i + 1
        t.insert(pos, n)

DEPOIS = {k: [c['s'], c['j']] + c['t'] for k, c in cells.items()}
TOTAL = sum(len(v) for v in DEPOIS.values())
cnt = collections.Counter(x for v in DEPOIS.values() for x in v)
# nenhum outro frasco muda de célula
for k in cells:
    a = [x for x in ANTES[k] if x not in REALOCA]; d = [x for x in DEPOIS[k] if x not in REALOCA]
    assert a == d, k
for n in cnt: assert cnt[n] <= CAP[MAPA[n]['t']], n

# ---- reescreve as linhas de célula
def cell_line(o, c):
    s = f"**{o}** — **Seguro: {c['s']}** ({c['sp']}). {c['sdesc']} — tier {c['st']}. **A jogada: {c['j']}** — {c['jr']} (tier {c['jt']})."
    if c['t']: s += ' **Também:** ' + ', '.join(c['t']) + '.'
    return s
for k, i in where.items():
    L[i] = cell_line(k[1], cells[k])
P = '\n'.join(L)

# ---- custo por uso (recalculado)
def frascos_faixa(b, dia):
    return len({x for (bb, o), v in DEPOIS.items() if bb == b and ((o in DIA_OCC) == dia) for x in v})
EMO = ['🔴', '🟠', '🟡', '🔵', '🟣', '⚫']
def fmt(x, d=1): return f'{x:.{d}f}'.replace('.', ',')
rows = []
FD = {}
for b in range(6):
    fd, fn = frascos_faixa(b, True), frascos_faixa(b, False); FD[b] = (fd, fn)
    rows.append(f"| {EMO[b]} {BANDS[b][1]} | {DIAS[b]} | {NOITES[b]} | {fd} / {fn} | {fmt(DIAS[b]/fd if fd else 0)} / {fmt(NOITES[b]/fn if fn else 0)} |")
i0 = P.index('| 🔴 Muito Quente | 27 | 1 |'); i1 = P.index('\n\n**Regras (vigentes desde 19/set/2026')
P = P[:i0] + '\n'.join(rows) + P[i1:]
P = rep1(P, 'na grade v59 (Dia = Lazer', 'na grade v60 (Dia = Lazer')
P = rep1(P, 'hoje 38 frascos disputam a noite de Ameno e 41 a de Fresco.', f'hoje {FD[2][1]} frascos disputam a noite de Ameno e {FD[3][1]} a de Fresco.')
P = rep1(P, 'Fresco de dia: 44 frascos para ~55 dias.', f'Fresco de dia: {FD[3][0]} frascos para ~55 dias.')

# ---- distribuição
GRADE = sorted(cnt)  # 71 com Bedouin Rose
tiers_g = collections.Counter(MAPA[n]['t'] for n in GRADE)
grp = collections.defaultdict(list)
for n in GRADE: grp[cnt[n]].append(n)
dist = [f"{len(GRADE)} frascos ativos · **{TOTAL} alocações** · média {fmt(TOTAL/len(GRADE), 2)} janelas.", '']
for k in sorted(grp, reverse=True):
    dist.append(f"- **{k} janelas ({len(grp[k])}):** " + ', '.join(sorted(grp[k], key=str.lower)))
i0 = P.index('69 frascos ativos · **651 alocações**'); i1 = P.index('\n\n**Destaques (Seguro + Jogada)')
P = P[:i0] + '\n'.join(dist) + P[i1:]
abaixo = sorted([n for n in GRADE if cnt[n] < CAP[MAPA[n]['t']]], key=lambda n: (TRANK[MAPA[n]['t']], n))

# ---- cabeçalho
i0 = P.index('# Playbook v59'); i1 = P.index('> **Fonte de tier:**')
mov = lambda n: ', '.join(f"{EMO[b]}{BANDS[b][1]}×{OCC_AB[o]}" for b, o in sorted(NOVAS[n]))
P = P[:i0] + f"""# Playbook v60 — {len(GRADE)} frascos na grade (Bedouin Rose entra como C)

> **v60 (23/out/2026) — regerado por script (`gerador_v60.py`).** Seguros, jogadas e sprays de Seguro **idênticos à v59** (regra de incumbência, decisão dele de 25/set/2026). Mudam só os *Também* de cinco frascos — nenhum outro frasco muda de célula (verificado por script):
> - **Blue Sapphire Supercharged** (S, frasco desde 09/out/2026) entra: {cnt['Blue Sapphire Supercharged']} janelas — {mov('Blue Sapphire Supercharged')}.
> - **Bedouin Rose** sobe de D para **C** e entra na grade (decisão dele, 23/out/2026): {cnt['Bedouin Rose']} janelas, só células-lar — {mov('Bedouin Rose')}.
> - **Favonius** sai do Trab. Formal e ganha Cozy Noite ({cnt['Favonius']} janelas, só noite Ameno–MF). **Omnia Omnibus Ubique** mantém as {cnt['Omnia Omnibus Ubique']} células de noite, com aptidões reordenadas. **Scarlet Sands** alarga o envelope até o Frio ({cnt['Scarlet Sands']} janelas). Envelopes e aptidões dos três foram decididos pelo Claude a pedido dele (23/out/2026), com base na pirâmide verificada + Fragrantica/Parfumo/site da marca — **continuam estimativa, não uso declarado**.
> - Desempate de *fit*: entre células de mesmo fit, fica a de faixa mais frequente no ano (normais INMET da tabela de custo por uso).
>
""" + P[i1:]
P = rep1(P, 'Os tiers da grade abaixo conferem com ela em 05/out/2026.', 'Os tiers da grade abaixo conferem com ela em 23/out/2026 (81 frascos próprios).')
P = rep1(P, 'Cap de 3 Seguros e 2 jogadas por frasco por faixa. Também = restante, ordenado por tier → aptidão → nome.',
         'Cap de 3 Seguros e 2 jogadas por frasco por faixa. Também = restante, ordenado por tier → aptidão → centralidade (frasco novo entra no fim do seu bloco tier × aptidão).')
i0 = P.index('> **Total: 651 alocações'); i1 = P.index('\n\n> **Princípio:**')
P = P[:i0] + (f"> **Total: {TOTAL} alocações · {len(GRADE)} frascos · " + ' · '.join(f"{t} {tiers_g[t]}" for t in 'SABC') + "** (D: nenhum frasco próprio). "
              "Abaixo do teto por falta de célula elegível: " + ', '.join(f"{n} ({MAPA[n]['t']}, {cnt[n]})" for n in abaixo) + ".\n"
              "> **Camada custo (10 frascos):** fora do playbook, do ranking e de qualquer análise — usada só em academia e corrida (regra dele, 25/set/2026); registrada no diário no slot **E** (Esporte). Empire entrou em 11/out/2026.") + P[i1:]

# ---- notas, histórico
P = rep1(P, '**Pendências que reabrem:** (1) onde ele de fato usa Favonius, Omnia Omnibus Ubique e Scarlet Sands (envelopes hoje estimados pelo Claude); (2) se ele quiser que os Seguros reflitam os novos S (hoje vale incumbência, por decisão dele).',
         '**Pendências que reabrem:** (1) uso real de Favonius, Omnia Omnibus Ubique, Scarlet Sands e Blue Sapphire Supercharged no diário — envelopes decididos pelo Claude em 23/out/2026, a pedido dele, sem uso declarado; a coluna NotaDia (acerto do dia, 1–5) é o dado que confirma ou derruba cada célula; (2) se ele quiser que os Seguros reflitam os novos S (hoje vale incumbência, por decisão dele).')
P = rep1(P, '| Versão | Data | O que mudou | Alocações |\n|---|---|---|---|\n',
         '| Versão | Data | O que mudou | Alocações |\n|---|---|---|---|\n'
         f'| v60 | 23/out/2026 | Entra Blue Sapphire Supercharged (S); Bedouin Rose D→C entra na grade; Favonius sai do Trab. Formal e ganha Cozy Noite; Scarlet Sands até o Frio; Omnia reordenado | {TOTAL} |\n')
assert 'v59 —' not in P.split('\n')[0]
open('out/playbook_v60.md', 'w', encoding='utf-8').write(P)
print('playbook v60:', TOTAL, 'alocações ·', len(GRADE), 'frascos ·', dict(tiers_g))
for n in REALOCA: print(' ', n, cnt[n], sorted((BANDS[b][1], o) for b, o in NOVAS[n]))
print(' abaixo do teto:', [(n, cnt[n]) for n in abaixo]); print(' FD', FD)

# ======== PARTE B — fichas_tecnicas_v10.md ========
ENV = ['🔴MQ', '🟠Q', '🟡A', '🔵F', '🟣Fr', '⚫MF']
def env(n): m = MAPA[n]; return f"{ENV[m['a']]}–{ENV[m['b']]}"
F = open('fichas_tecnicas_v9.md', encoding='utf-8').read()
FL = F.split('\n')
ROW = re.compile(r'^\| (\d+) \| \*\*(.+?)\*\* \|')
rows = {}; order = []
for i, l in enumerate(FL):
    m = ROW.match(l)
    if m:
        cols = l.split(' | '); rows[m.group(2)] = (i, l); order.append(m.group(2))
assert len(rows) == 79, len(rows)

def cols_of(l):
    c = [x for x in l.strip().strip('|').split(' | ')]
    c = [x.strip() for x in c]; assert len(c) == 17, (len(c), l[:60]); return c
def join(c): return '| ' + ' | '.join(c) + ' |'

LEIT = {
 'Favonius': "Rosa 'cristalina' com incenso do topo à base, lado herbal-verde (artemísia, clary sage, gerânio) e cypriol seco sobre oud-patchouli-cedro: a rosa-oud mais austera e menos doce da grade. **O frasco já era do Guilherme e constava como Amostra (tier A) no tabelão — erro de cadastro corrigido em 23/set/2026; tier S dado por ele.** **v10 (23/out/2026): sai do Trab. Formal e fica só na noite (N.Inf · N.For · Cozy N, Ameno–MF)** — decisão do Claude a pedido dele: extrait de projeção e sillage altos (Fragrantica: 'Woody Chypre', resenhas de ambiente cheio; Parfumo: longevidade 8,5 · sillage 8,0) não é perfume de escritório; o diário dele confirma noite (09 e 15/out/2026). **Continua estimativa, não uso declarado.** **Registro franco:** o Rosal tem 7 frascos (6 S + Bedouin Rose C); o que separa o Favonius é o eixo herbal-incenso (fronteira Monge), não o território.",
 'Omnia Omnibus Ubique': "Manteiga de íris empoada ancorada num oud floral macio, com açafrão e jasmim; base baunilha-sândalo-musk cremoso. **Tier S dado pelo Guilherme (23/set/2026).** **v10 (23/out/2026): aptidões reordenadas — N.For 3 · Cozy N 3 · N.Inf 2** (eram N.Inf 3 · N.For 2 · Cozy N 2); mesmas 12 células de noite. Base: Fragrantica (Oriental Woody; acordes empoado-baunilha-amadeirado-oud-íris), Parfumo (cremoso-amadeirado; sillage 7,6 — moderado) e a escala dia/noite 3/4 do site da Memo. Decisão do Claude a pedido dele — **estimativa, não uso declarado**. Fronteira Sultão (açafrão-oud) registrada; toca o gap #4 (sândalo cremoso) só como base, não fecha.",
 'Scarlet Sands': "Framboesa e mandarina sobre flor de laranjeira, jasmim e frésia; akigalawood, baunilha e âmbar na base — floral branco frutado-doce. **Tier S dado pelo Guilherme (23/set/2026).** **v10 (23/out/2026): envelope alarga até o Frio (Quente–Frio, 12 janelas)** — topo cítrico-frutado aguenta calor e a base âmbar-baunilha sustenta noite fria (Fragrantica: Oriental Floral; Parfumo: um resenhista o descreve para dias quentes que viram noites frescas). É o único dos frascos revistos em 23/out com janela de **Dia** (Lazer) — e o gargalo da coleção é o slot Dia. Decisão do Claude a pedido dele — **estimativa, não uso declarado**. Fica fora do Muito Quente: doçura 2 no mormaço. Resenhas divergem sobre performance (de 'powerhouse' a 'queria mais força') — calibrar sprays pelo diário.",
 'Bedouin Rose': "Rosa oriental resinosa (tâmara-canela-açafrão sobre rosa Taif-imortelle, base âmbar-olíbano-oud). **v10 (23/out/2026): tier C dado por ele (aba Tiers, 14/out) e entrada na grade por decisão dele** — 'os C ficam'. Como C, só nas células-lar: Cozy Noite no Frio e no Muito Frio (2 janelas). Frasco de 10 ml experimental — fora de custo por uso e de desbaste por volume. Vizinho direto do Brompton.",
}
NOVA_BSS = ['', '**Blue Sapphire Supercharged**', 'Boadicea', "Parfum (marca: 'pure parfum'; varejistas listam EDP)", 'Rosal', 'S',
 "laranja-limão (aldeídos) / rosa-cardamomo-jasmim (framboesa e ylang citados pela marca) / oud-âmbar-patchouli, toque de couro (2022 — pirâmide: Selfridges e Neiman Marcus; descrição oficial Boadicea; verificada em 23/out/2026)",
 'S', '1', '4', 'S', env('Blue Sapphire Supercharged'), FICHA['Blue Sapphire Supercharged'][2], '3', str(cnt['Blue Sapphire Supercharged']),
 'Favonius, Radical Rose, Les Sables Roses, Brompton (rosa-oud); Blue Sapphire (avaliado, A — a versão original)',
 "Rosa-oud de saída cítrica-aldeídica e coração framboesa-cardamomo: o rosa-oud mais luminoso da grade, papel de **sair** (N.For e N.Inf 3), não de ficar em casa. **Tier S dado por ele (frasco desde 09/out/2026; usado em 19/out à noite, 14 °C).** Fragrantica: Oriental Floral; resenhas relatam 12h+ e versatilidade no ano todo. Envelope Ameno–MF, aptidões e sprays derivados pelo Claude em 23/out/2026 — **estimativa, não uso declarado**. **Registro franco:** com ele o Rosal vai a 7 frascos (6 S + Bedouin Rose C) e quatro rosa-oud S disputam as mesmas noites Ameno–MF — compra futura de rosa-oriental parte de redundância alta no gate 3."]
NOVA_EMP = ['', '**Empire**', 'Mykonos', 'n/d', 'Marinho', 'A',
 "lima-bergamota-mandarina-frutas tropicais / laranja siciliana-gengibre-jasmim-água de coco / patchouli-musk branco-tonka (2026, Mykonos x RRQ — pirâmide de varejista, arabicparfums.com; Parfumo o descreve como híbrido de Afternoon Swim e Virgin Island Water)",
 'F', '1', '2', 'C', '🔴MQ–🟡A', FICHA['Empire'][2], '4', '—', 'Aquatica, Island Dreams', "Cítrico-tropical de calor. Camada custo (tier A dado por ele, 11/out/2026): corrida, academia e praia pós-corrida — fora de qualquer análise."]

grade_rows, custo_rows = [], []
for n in order:
    i, l = rows[n]; c = cols_of(l)
    if n in MUDA:
        c[5] = MAPA[n]['t']; c[11] = env(n); c[12] = FICHA[n][2]; c[16] = LEIT[n]
    if n in cnt:
        c[14] = str(cnt[n])
    elif n == 'Bedouin Rose':
        raise AssertionError
    (custo_rows if MAPA[n]['c'] else grade_rows).append(c)
    if n not in REALOCA and n in cnt: assert int(cols_of(l)[14]) == cnt[n], n
grade_rows.append(NOVA_BSS); custo_rows.append(NOVA_EMP)
grade_rows.sort(key=lambda c: c[1].strip('*').lower())
for k, c in enumerate(grade_rows + custo_rows, 1): c[0] = str(k)
assert len(grade_rows) == 71 and len(custo_rows) == 10
# reconstrói o arquivo: substitui as linhas de tabela no lugar
first_g = min(rows[n][0] for n in order if not MAPA[n]['c']); last_g = max(rows[n][0] for n in order if not MAPA[n]['c'])
first_c = min(rows[n][0] for n in order if MAPA[n]['c']); last_c = max(rows[n][0] for n in order if MAPA[n]['c'])
FL = FL[:first_g] + [join(c) for c in grade_rows] + FL[last_g + 1:first_c] + [join(c) for c in custo_rows] + FL[last_c + 1:]
F = '\n'.join(FL)

# cabeçalho
i0 = F.index('# Fichas Técnicas — v9'); i1 = F.index('> **Fonte de tier:**')
F = F[:i0] + """# Fichas Técnicas — v10 (23/out/2026) · grade 71 (Bedouin Rose entra como C) + camada custo 10

> **v10 — regerado por `gerador_v60.py`.** Entram **Blue Sapphire Supercharged** (grade, S) e **Empire** (camada custo, A). **Bedouin Rose** D→C entra na grade (2 janelas, células-lar). Revistos por decisão do Claude a pedido dele (23/out/2026), com Fragrantica/Parfumo/site da marca: **Favonius** (sai do Trab. Formal, ganha Cozy Noite), **Omnia Omnibus Ubique** (aptidões reordenadas), **Scarlet Sands** (envelope até o Frio). Nenhum outro campo de nenhum outro frasco muda. Janelas = playbook v60 (668 alocações).
>
""" + F[i1:]
F = rep1(F, 'Tiers desta ficha conferem com ela em 05/out/2026.', 'Tiers desta ficha conferem com ela em 23/out/2026.')
F = rep1(F, '## Grade — 69 frascos + Bedouin Rose (D, 0 janelas) (ordem alfabética)', '## Grade — 71 frascos (ordem alfabética)')
F = rep1(F, '## Camada custo — 9 frascos (só academia e corrida)', '## Camada custo — 10 frascos (só academia e corrida · slot E no diário)')
F = rep1(F, '| **Janelas** | nº de células (playbook v59) |', '| **Janelas** | nº de células (playbook v60) |')
F = rep1(F, 'Vibrato só dia; Tygar Extrait noite de calor; Les Sables Roses noites exceto as quentes)',
         'Vibrato só dia; Tygar Extrait noite de calor; Les Sables Roses noites exceto as quentes; Bedouin Rose C na grade)')

# leituras transversais — recalculadas da grade
G = [n for n in cnt]
ex = collections.Counter(FICHA[n][1] for n in G); dc = collections.Counter(MAPA[n]['d'] for n in G)
pc = collections.Counter(MAPA[n]['p'] for n in G); tc = collections.Counter(MAPA[n]['td'] for n in G)
comp = f"""| Dimensão | Distribuição | Leitura |
|---|---|---|
| Eixo DNA | P {ex['P']} · S {ex['S']} · F {ex['F']} | O eixo primário declarado tem {ex['P']} frascos; o secundário ambarado-doce, {ex['S']} (Blue Sapphire Supercharged entra); {ex['F']} são variedade deliberada |
| Doçura | 0: {dc[0]} · 1: {dc[1]} · 2: {dc[2]} · 3: {dc[3]} | {dc[2]+dc[3]} frascos com baunilha/tonka/benjoim/caramelo protagonizando a base. Desde 18/out/2026 doçura não é demérito em avaliação de compra (fala dele) |
| Peso | 1: {pc[1]} · 2: {pc[2]} · 3: {pc[3]} · 4: {pc[4]} · 5: {pc[5]} | {pc[4]+pc[5]} frascos de peso 4–5 disputam as ~300 noites Ameno/Fresco; {pc[1]+pc[2]} de peso 1–2 servem o slot Dia de calor |
| Td | ↑carregado {tc['C']} · ↑seco {tc['S']} · neutro {tc['N']} | Num dia de mormaço (Td ≥20) {tc['C']} frascos sobem e {tc['S']} descem; em ar seco de inverno inverte |
"""
i0 = F.index('| Dimensão | Distribuição | Leitura |'); i1 = F.index('\n### 2. Mapa de Td')
F = F[:i0] + comp + F[i1:]
lst = lambda k: ' · '.join(sorted([n for n in G if MAPA[n]['td'] == k], key=str.lower))
td = f"""### 2. Mapa de Td por frasco (substitui as listas genéricas do Passo 5 da skill)

- **Sobem no ar carregado (Td ≥ 20°C) · descem no ar seco:** {lst('C')}
- **Sobem no ar seco (Td < 12°C) · descem no mormaço:** {lst('S')}
- **Neutros — ordem padrão do playbook:** {lst('N')}

### 3. Territórios de redundância (contagem por vizinhança declarada nas fichas)
"""
i0 = F.index('### 2. Mapa de Td'); i1 = F.index('\n| Território | n | Frascos | Leitura |')
F = F[:i0] + td + F[i1:]
F = rep1(F, "| Rosa oriental | 8 | Radical Rose, Odéon, Brompton, Bedouin Rose, Oud Satin Mood, Serenity, Les Sables Roses, Favonius |",
            "| Rosa oriental | 9 | Radical Rose, Odéon, Brompton, Bedouin Rose, Oud Satin Mood, Serenity, Les Sables Roses, Favonius, Blue Sapphire Supercharged |")
F = rep1(F, "**6 S em 8 (Radical Rose, Odéon, Brompton, Oud Satin Mood, Les Sables Roses, Favonius) após a calibração de 25/set — o território mais denso de S da grade.** Favonius é o único seco-herbal-incenso do grupo",
            "**7 S em 9 (Radical Rose, Odéon, Brompton, Oud Satin Mood, Les Sables Roses, Favonius, Blue Sapphire Supercharged) — o território mais denso de S da grade.** Favonius é o único seco-herbal-incenso do grupo; Blue Sapphire Supercharged, o mais luminoso (cítrico-aldeído-framboesa). Próxima rosa-oriental parte de redundância alta no gate 3")
F = rep1(F, "| Oud nomeado | 15 | Ombre Nomade, Oud Satin Mood, Monarch, Al Qurashi Blend, Oudós Lux Solis, Oud Kyphi, Liwa, Royal Oud, Carved Oud, African Leather, Imperial, Serenity, Les Sables Roses, Favonius, Omnia Omnibus Ubique | 15 frascos",
            "| Oud nomeado | 17 | Ombre Nomade, Oud Satin Mood, Monarch, Al Qurashi Blend, Oudós Lux Solis, Oud Kyphi, Liwa, Royal Oud, Carved Oud, African Leather, Imperial, Serenity, Les Sables Roses, Favonius, Omnia Omnibus Ubique, Bedouin Rose, Blue Sapphire Supercharged | 17 frascos")
# pendências
i0 = F.index('### 5. Pendências (05/out/2026)')
F = F[:i0] + """### 4. Pendências (23/out/2026)

- **Quatro frascos com envelope e aptidão decididos pelo Claude, a pedido dele, sem uso declarado** (Favonius, Omnia Omnibus Ubique, Scarlet Sands, Blue Sapphire Supercharged — 23/out/2026). O que confirma ou derruba cada célula: o diário, com a coluna NotaDia (acerto do dia, 1–5).
- **Tygar Extrait:** se render em noite fresca, o envelope alarga ao Fresco e a grade regera (S comporta 12; hoje tem 6).
- **Arquétipos:** a divisão Mediterrâneo + Solar (Imagination, Lucius, Tygar Extrait), pedida em 05/out/2026, ainda não está propagada às fichas, ao mapa e ao observatório.
"""
assert 'v59' not in F.replace('playbook v59', '').split('## Como ler')[0]
open('out/fichas_tecnicas_v10.md', 'w', encoding='utf-8').write(F)
print('fichas v10:', len(grade_rows), 'grade +', len(custo_rows), 'custo · eixo', dict(ex), '· td', dict(tc))

# ======== PARTE C — frascos_v8.py + modelo_diario_v1_7.py ========
S = open('frascos_v7.py', encoding='utf-8').read()
ordem = [m['n'] for m in fv7.MAPA]
grade_ord = sorted([n for n in MAPA if not MAPA[n]['c']], key=str.lower)
custo_ord = [n for n in ordem if MAPA[n]['c']] + ['Empire']
KEYS = ['n', 'casa', 'arq', 't', 'td', 'p', 'd', 'a', 'b', 'c', 'r', 'u']
def mline(n):
    m = MAPA[n]; return ' {' + ','.join(f'"{k}":' + json.dumps(m[k], ensure_ascii=False) for k in KEYS) + '},'
mapa_txt = 'MAPA = [\n' + '\n'.join(mline(n) for n in grade_ord + custo_ord) + '\n]'
i0 = S.index('MAPA = ['); i1 = S.index(']\n\n# FICHA') + 1
S = S[:i0] + mapa_txt + S[i1:]
def fline(n):
    conc, eixo, apt, sp, _ = FICHA[n]
    jan = 'None' if MAPA[n]['c'] else str(cnt[n])
    com = {'Vibrato': '   # sai das noites por decisão dele (05/out/2026)',
           'Bedouin Rose': '   # C na grade por decisão dele (23/out/2026): só células-lar',
           'Favonius': '   # v8: envelope/aptidão revistos pelo Claude a pedido dele (23/out/2026) — estimativa',
           'Omnia Omnibus Ubique': '   # v8: idem', 'Scarlet Sands': '   # v8: idem — envelope até o Frio',
           'Blue Sapphire Supercharged': '   # v8: frasco desde 09/out/2026; envelope/aptidão do Claude — estimativa'}.get(n, '')
    return f' {json.dumps(n, ensure_ascii=False)}: ({json.dumps(conc, ensure_ascii=False)}, "{eixo}", "{apt}", {sp}, {jan}),{com}'
ficha_txt = 'FICHA = {\n' + '\n'.join(fline(n) for n in grade_ord) + '\n # camada custo (janelas = None: fora da grade)\n' + '\n'.join(fline(n) for n in custo_ord) + '\n}'
i0 = S.index('FICHA = {'); i1 = S.index('}\n\n# NOVO') + 1
S = S[:i0] + ficha_txt + S[i1:]
S = rep1(S, ' "Brompton": (12, 4, "F"),          # extrait; "começar em 2 sprays"\n',
         ' "Brompton": (12, 4, "F"),          # extrait; "começar em 2 sprays"\n "Blue Sapphire Supercharged": (10, 4, "F"),  # Fragrantica: relatos de 12h+ — estimativa\n')
# validações
S = rep1(S, '    assert len(frascos) == 79, len(frascos)', '    assert len(frascos) == 81, len(frascos)')
S = rep1(S, '    assert len(grade) == 70', '    assert len(grade) == 71')
tcg = collections.Counter(MAPA[n]['td'] for n in grade_ord)
S = rep1(S, '    assert td_cnt == {"C": 21, "S": 25, "N": 24}, td_cnt', f'    assert td_cnt == {{"C": {tcg["C"]}, "S": {tcg["S"]}, "N": {tcg["N"]}}}, td_cnt')
S = rep1(S, '== 651, sum(', f'== {TOTAL}, sum(')
tall = collections.Counter(MAPA[n]['t'] for n in MAPA)
S = rep1(S, '    assert tiers == {"S": 33, "A": 30, "B": 13, "C": 2, "D": 1}, tiers', f'    assert tiers == {{"S": {tall["S"]}, "A": {tall["A"]}, "B": {tall["B"]}, "C": {tall["C"]}, "D": 0}}, tiers')
S = rep1(S, 'print("OK — 79 frascos · Td grade C21/S25/N24 · janelas 651 · tiers"', f'print("OK — 81 frascos · Td grade C{tcg["C"]}/S{tcg["S"]}/N{tcg["N"]} · janelas {TOTAL} · tiers"')
i0 = S.index('"""\nfrascos_v7.py'); i1 = S.index('Fontes, por campo:')
S = S[:i0] + f'''"""
frascos_v8.py — base de dados dos 81 frascos para o modelo diário (v8, 23/out/2026 — gerado por gerador_v60.py)

v8: entram Blue Sapphire Supercharged (grade, S) e Empire (camada custo, A); Bedouin Rose D→C entra na grade
(só células-lar); Favonius, Omnia Omnibus Ubique e Scarlet Sands com envelope/aptidão revistos pelo Claude a pedido
dele (23/out/2026 — estimativa, não uso declarado). Tiers conferidos com a aba Tiers do Google Sheets "Diário de Uso —
Perfumes" em 23/out/2026. Janelas = playbook v60 ({TOTAL}). A camada custo (10) fica aqui só para o modelo sugerir
quando a entrada traz "Extra: corrida" ou "Extra: academia" — fora de qualquer análise da coleção (regra dele, 25/set/2026).

''' + S[i1:]
S = S.replace('= transcritos das fichas_tecnicas_v9.md (grade 69 + Bedouin Rose + camada custo 9).', '= transcritos das fichas_tecnicas_v10.md (grade 71 + camada custo 10).')
S = S.replace('Validação embutida: 79 frascos · Td grade C21/S25/N24 · janelas = 651.', f'Validação embutida: 81 frascos · Td grade C{tcg["C"]}/S{tcg["S"]}/N{tcg["N"]} · janelas = {TOTAL}.')
open('out/frascos_v8.py', 'w', encoding='utf-8').write(S)

M = open('modelo_diario_v1_6.py', encoding='utf-8').read()
M = rep1(M, 'modelo_diario_v1_6.py — Top 10 do dia em DOIS SLOTS (v1.6, 05/out/2026 — base frascos_v7)',
         'modelo_diario_v1_7.py — Top 10 do dia em DOIS SLOTS (v1.7, 23/out/2026 — base frascos_v8)\n\n'
         'v1.7 vs v1.6: (1) base frascos_v8 (grade 71: entram Blue Sapphire Supercharged e Bedouin Rose C; Empire na camada\n'
         'custo); (2) linhas do diário com slot E (Esporte — camada custo, corrida/academia) são lidas: "(esporte)" ou "(e)"\n'
         'no fim do nome é removido como "(manhã)"/"(noite)". Pesos, filtros e ranking idênticos à v1.6.')
M = rep1(M, 'from frascos_v7 import carregar', 'from frascos_v8 import carregar')
M = rep1(M, r'''r"\s*\((manh[ãa]|noite|dia|m|n)\)\s*$"''', r'''r"\s*\((manh[ãa]|noite|dia|esporte|m|n|e)\)\s*$"''')
M = M.replace('python3 modelo_diario_v1_6.py', 'python3 modelo_diario_v1_7.py')
M = rep1(M, '_Modelo v1.6 (dois slots · base frascos_v7, grade 69)', '_Modelo v1.7 (dois slots · base frascos_v8, grade 71)')
open('out/modelo_diario_v1_7.py', 'w', encoding='utf-8').write(M)
print('frascos_v8 + modelo v1.7 escritos')

# ======== PARTE D — tabelao_v2_8.md + .xlsx (espelho da aba Tiers em 23/out/2026) ========
import openpyxl, copy
T7 = open('tabelao_v2_7.md', encoding='utf-8').read(); L7 = T7.split('\n')
old = []
for l in L7:
    m = re.match(r'^\| (\d+) \| (.+?) \| (.+?) \| (.+?) \| (.*?) \| (.+?) \|$', l)
    if m: old.append(m.groups())
assert len(old) == 336
RENOME = {('amouage', 'guidance 46'): 'guidance 46 (exceptional extrait)', ('amouage', 'interlude 53'): 'interlude 53 (exceptional extrait)',
          ('amouage', 'jubilation 40'): 'jubilation 40 (exceptional extrait)', ('amouage', 'sequence'): 'sequence (essences)',
          ('initio', 'wild rush'): 'wild rush (2026)'}
ko = {(o[1].lower(), o[2].lower()): o for o in old}
rows = sorted(TIERS, key=lambda r: (r[0].lower(), r[1].lower()))
out, mud = [], []
seen = set()
for r in rows:
    casa, perf, _, posse, tier, data, obs = r
    k = (casa.lower(), perf.lower()); k7 = (k[0], RENOME.get(k, k[1])); seen.add(k7)
    o = ko.get(k7)
    v1 = o[5] if o else '—'
    out.append([casa, perf, posse, tier, v1])
    if not o: mud.append((casa, perf, f'Linha nova · {posse or "posse vazia"} · tier {tier or "pendente"}' + (f' ({data})' if data else '')))
    else:
        ch = []
        if o[2] != perf: ch.append(f'nome normalizado como na aba Tiers (era "{o[2]}")')
        if o[3] != posse: ch.append(f'posse {o[3]} → {posse}')
        if o[4] != tier: ch.append(f'tier {o[4] or "—"} → {tier or "pendente"}')
        if ch: mud.append((casa, perf, '; '.join(ch)))
for k, o in ko.items():
    if k not in seen: mud.append((o[1], o[2], 'Linha removida (duplicata — decisão dele, 23/out/2026)'))
assert len(out) == 355
POS = ['Frasco Gui', 'Frasco Bia', 'Amostra Bia', 'Amostra', 'Avaliado']
assert all(r[2] in POS for r in out)
dist = collections.Counter((r[3], r[2]) for r in out)
def drow(t):
    v = [dist[(t, p)] for p in POS]; return f"| {t or 'sem tier'} | {sum(v)} | " + ' | '.join(map(str, v)) + ' |'
fg = collections.Counter(r[3] for r in out if r[2] == 'Frasco Gui')
md = [f"# Tabelão v2.8 — S/A/B/C/D (23/out/2026)", '',
 '> **Espelho da aba *Tiers* da planilha Google Sheets "Diário de Uso — Perfumes"**, que é a fonte única de tier desde 04/out/2026 (decisão dele). Em conflito, a planilha vence; este arquivo é regerado a partir dela. Âncora: **S** compraria de novo amanhã · **A** alegria real · **B** respeito, cumpre função · **C** neutro · **D** fora (não comprar).',
 '>',
 f'> **v2.8 (23/out/2026):** limpeza da aba Tiers feita com ele — zero "a confirmar", zero posse vazia; duplicatas de Bracken e Elsewhere removidas; Bedouin Rose C; entram Blue Sapphire Supercharged (S) e Empire (A) como Frasco Gui e 20 linhas avaliadas desde 05/out. Um tier pendente por decisão dele (Amouage Enclave). {len(out)} linhas. Histórico na aba *Histórico* do xlsx.',
 '', '## Distribuição', '', '| Tier | Total | ' + ' | '.join(POS) + ' |', '|---|---|' + '---|' * len(POS)]
md += [drow(t) for t in ['S', 'A', 'B', 'C', 'D', '']]
md += [f"| **Total** | **{len(out)}** | " + ' | '.join(f"**{sum(1 for r in out if r[2] == p)}**" for p in POS) + ' |', '',
       f"Frasco Gui = {sum(fg.values())} (71 grade + 10 camada custo) · por tier: " + ' · '.join(f"{t} {fg[t]}" for t in 'SABCD') + '.', '',
       '## Tabelão', '', '| # | Casa | Perfume | Posse | Tier | # v1 |', '|---|---|---|---|---|---|']
md += [f"| {i} | {r[0]} | {r[1]} | {r[2]} | {r[3] or '—'} | {r[4]} |" for i, r in enumerate(out, 1)]
md += ['', '## Mudanças v2.7 → v2.8 (23/out/2026)', '', '| Casa | Perfume | Mudança |', '|---|---|---|']
md += [f"| {c} | {p} | {m} |" for c, p, m in mud]
# seção "Novos" da v2.7 mantida, com a frase de posse atualizada
i0 = T7.index('## Novos no tabelão — 17'); i1 = T7.index('## Tier D — detalhe')
novos = T7[i0:i1]
novos = re.sub(r'Posse de todos os 17 = \*\*a confirmar\*\*[^\n]*', 'Posse dos 17 resolvida com ele em 23/out/2026 (ver tabela acima).', novos)
md += ['', novos.rstrip(), '']
dd = [r for r in out if r[3] == 'D']
md += [f'## Tier D — detalhe ({len(dd)}, já incluídos na tabela)', '', '| Casa | Perfume | Posse |', '|---|---|---|'] + [f"| {r[0]} | {r[1]} | {r[2]} |" for r in dd]
open('out/tabelao_v2_8.md', 'w', encoding='utf-8').write('\n'.join(md) + '\n')

# ---- xlsx: mesma estrutura e formatação da v2.7
wb = openpyxl.load_workbook('tabelao_v2_7.xlsx')
ws = wb['Tabelão v2']
st = [copy.copy(ws.cell(2, c)._style) for c in range(1, 7)]
ws.delete_rows(2, ws.max_row)
for i, r in enumerate(out, 2):
    for c, v in enumerate([i - 1, r[0], r[1], r[2], r[3] or None, r[4]], 1):
        cell = ws.cell(i, c, v); cell._style = copy.copy(st[c - 1])
last = len(out) + 1
ws.auto_filter.ref = f'A1:F{last}'
wd = wb['Distribuição']
for row in wd.iter_rows():
    for cell in row:
        if isinstance(cell.value, str) and cell.value.startswith('='):
            cell.value = cell.value.replace('$337', f'${last}').replace('"a confirmar*"', '"a confirmar*"')
for row in wd.iter_rows():
    for cell in row:
        if isinstance(cell.value, str) and cell.value.startswith('Fonte: aba Tiers'):
            cell.value = 'Fonte: aba Tiers da planilha Google Sheets "Diário de Uso — Perfumes" (fonte única de tier desde 04/out/2026), espelhada em 23/out/2026. A coluna "A confirmar" ficou zerada na limpeza de 23/out/2026.'
wh = wb['Histórico']
hs = [copy.copy(wh.cell(2, c)._style) for c in range(1, 6)]
wh.insert_rows(2, len(mud))
for i, (c, p, m) in enumerate(mud, 2):
    for j, v in enumerate(['v2.8', '23/out/2026', c, p, m], 1):
        cell = wh.cell(i, j, v); cell._style = copy.copy(hs[j - 1])
wdd = wb['Tier D (detalhe)']
ds = [copy.copy(wdd.cell(2, c)._style) for c in range(1, 4)]
wdd.delete_rows(2, wdd.max_row)
for i, r in enumerate(dd, 2):
    for j, v in enumerate([r[0], r[1], r[2]], 1):
        cell = wdd.cell(i, j, v); cell._style = copy.copy(ds[j - 1])
wb.save('out/tabelao_v2_8.xlsx')
print('tabelão v2.8:', len(out), 'linhas ·', len(mud), 'mudanças · FG', dict(fg))

# ======== PARTE E — mapa_climatico_colecao_v6.html · observatorio_colecao_v4_8.html · avaliacao_perfumes_v52 ========
KEYS_M = ['n', 'casa', 'arq', 't', 'td', 'p', 'd', 'a', 'b', 'c', 'r', 'u']
# ---------------- mapa climático
H = open('mapa_climatico_colecao_v5.html', encoding='utf-8').read()
i0 = H.index('const DATA=') + len('const DATA='); i1 = H.index('];', i0) + 1
D = json.loads(H[i0:i1]); old = {d['n']: d for d in D}
NEW = []
for n in sorted([n for n in MAPA if not MAPA[n]['c']], key=str.lower):
    d = dict(old.get(n, {'s': 'B. Sapphire SC' if n == 'Blue Sapphire Supercharged' else n}))
    for k in KEYS_M: d[k] = MAPA[n][k]
    NEW.append(d)
assert len(NEW) == 71
H = H[:i0] + json.dumps(NEW, ensure_ascii=False, separators=(',', ':')) + H[i1:]
H = rep1(H, 'Mapa climático da coleção v5 — temperatura × ponto de orvalho × chuva · 70 frascos · fichas v9 · 05/out/2026',
         'Mapa climático da coleção v6 — temperatura × ponto de orvalho × chuva · 71 frascos · fichas v10 · 23/out/2026')
H = rep1(H, '<p>70 frascos posicionados', '<p>71 frascos posicionados')
j0 = H.index('v5 de 05/out/2026:'); j1 = H.index('Sem ocasi', j0)
H = H[:j0] + ('v6 de 23/out/2026: grade de 71 frascos (fichas v10) — entra Blue Sapphire Supercharged (S); Bedouin Rose passa a C, dentro da grade; '
              'Scarlet Sands alarga o envelope até o Frio; Favonius e Omnia Omnibus Ubique mudam só de ocasião (não muda posição). Camada custo fora do mapa '
              '(usada só em academia e corrida — regra de 25/set/2026). Tiers conferidos com a aba Tiers do Google Sheets. Índice de chuva do Blue Sapphire Supercharged: hipótese do Claude pela pirâmide, como os demais. ') + H[j1:]
H = rep1(H, 'aria-label="Dispersão dos 70 frascos', 'aria-label="Dispersão dos 71 frascos')
open('out/mapa_climatico_colecao_v6.html', 'w', encoding='utf-8').write(H)

# ---------------- observatório
O = open('observatorio_colecao_v4_6.html', encoding='utf-8').read()
m = re.search(r'const FR = (\[.*?\]);', O); FR = json.loads(m.group(1))
TV = {'S': 10, 'A': 8, 'B': 6, 'C': 4, 'D': 2}
for f in FR:
    if f['nome'] == 'Bedouin Rose': f['tier'] = 'C'; f['tv'] = 4; f.pop('out', None)
    if f['nome'] == 'Scarlet Sands': f['frio'] = 6
FR.append({'casa': 'Boadicea', 'nome': 'Blue Sapphire Supercharged', 'calor': 3, 'ameno': 7, 'frio': 9, 'arq': 'O Rosal', 'tier': 'S', 'tv': 10})
FR.sort(key=lambda f: (f['casa'].lower(), f['nome'].lower()))
assert len(FR) == 71
O = O[:m.start(1)] + json.dumps(FR, ensure_ascii=False) + O[m.end(1):]
i = O.index('const PB = {'); j = O.index('\n};', i)
PB = json.loads(O[i + len('const PB = '):j + 2].rstrip(';'))
BK = ['mq', 'q', 'am', 'fr', 'fri', 'mf']; OK = {'Lazer': 'lazer', 'Trab. Informal': 'ti', 'Trab. Formal': 'tf', 'Cozy Dia': 'cd',
      'Noite Informal': 'ni', 'Noite Formal': 'nf', 'Cozy Noite': 'cn'}
n_ok = 0
for (b, o), c in cells.items():
    old_c = PB[BK[b]][OK[o]]
    assert old_c['s'][0] == c['s'] and old_c['j'][0] == c['j'], (b, o)
    old_c['t'] = c['t']; n_ok += 1
assert n_ok == 38
O = O[:i] + 'const PB = ' + json.dumps(PB, ensure_ascii=False, indent=0) + O[j + 2:]
O = rep1(O, 'Observatório da Coleção v4.6 — tiers S/A/B/C/D · 69 na grade · 16 arquétipos · playbook v59',
         'Observatório da Coleção v4.8 — tiers S/A/B/C · 71 na grade · 16 arquétipos · playbook v60')
O = rep1(O, 'v4.6 · 05/out/2026 (playbook v59 · fichas v9 · tabelão v2.7) · <b>69 na grade</b> + Bedouin Rose fora da grade (tier D)',
         'v4.8 · 23/out/2026 (playbook v60 · fichas v10 · tabelão v2.8) · <b>71 na grade</b> (Blue Sapphire Supercharged entra; Bedouin Rose C, dentro da grade)')
O = rep1(O, 'com a grade de 69 frascos, camada custo de 10', 'com a grade de 71 frascos')
O = rep1(O, '/* ============ DADOS — 05/out/2026: 70 frascos (grade 69 + Bedouin Rose D); camada custo fora */', '/* ============ DADOS — 23/out/2026: 71 frascos (grade v60); camada custo fora */')
O = rep1(O, f'/* ============ PLAYBOOK v59 — grade literal (651 alocações · 69 frascos ·', f'/* ============ PLAYBOOK v60 — grade literal ({TOTAL} alocações · 71 frascos ·')
O = O.replace('célula literal do playbook v59', 'célula literal do playbook v60')
open('out/observatorio_colecao_v4_8.html', 'w', encoding='utf-8').write(O)

# ---------------- avaliação
A = open('avaliacao_perfumes_v51_2026-10-05.html', encoding='utf-8').read()
k0 = A.index('/*<<D>>*/') + len('/*<<D>>*/'); k1 = A.index('/*<</D>>*/')
DA = json.loads(A[k0:k1])
TN = {'S': 5, 'A': 4, 'B': 3, 'C': 2, 'D': 1, '': 0}
def norm(x): return re.sub(r'\s*\(.*?\)\s*', ' ', x).strip().lower()
idx = collections.defaultdict(list)
for r in DA: idx[(r[0].lower(), r[1].lower())].append(r); idx[('*', norm(r[1]), r[0].lower())].append(r)
ch, miss = [], []
for casa, perf, _, posse, tier, _, _ in TIERS:
    hits = idx.get((casa.lower(), perf.lower())) or idx.get(('*', norm(perf), casa.lower()))
    if not hits or len(hits) != 1: miss.append(f'{casa} | {perf}'); continue
    r = hits[0]; ps = '' if posse == 'Avaliado' else posse
    if r[2] != tier or r[21] != ps:
        ch.append((r[0], r[1], f"{r[2] or '—'}/{r[21] or 'Avaliado'} → {tier or '—'}/{posse}")); r[2] = tier; r[3] = TN[tier]; r[21] = ps
        if ps == 'Frasco Gui': r[9] = 'TEM'
UPD = {
 'Favonius': 'Rosa-incenso seca | Noite Ameno–MF (N.Inf · N.For · Cozy N) | v52 (23/out/2026): sai do Trab. Formal — envelope/aptidão decididos pelo Claude a pedido dele (Fragrantica/Parfumo), estimativa. Frasco ativo; tier S dado por ele.',
 'Omnia Omnibus Ubique': None, 'Scarlet Sands': 'Floral branco frutado | Lazer + Noite Quente–Frio | v52 (23/out/2026): envelope alarga até o Frio — decisão do Claude a pedido dele (Fragrantica/Parfumo), estimativa. Frasco ativo; tier S dado por ele.',
 'Bedouin Rose': 'Rosa oriental resinosa | Cozy Noite Frio–MF | v52 (23/out/2026): tier C dado por ele; entra na grade por decisão dele (só células-lar, 2 janelas). Frasco de 10 ml experimental.'}
for r in DA:
    if r[1] in UPD and UPD[r[1]] and r[21] == 'Frasco Gui': r[15] = UPD[r[1]]
names = {(r[0], r[1]) for r in DA}
NOVAS_L = []
if ('Boadicea', 'Blue Sapphire Supercharged') not in names:
    NOVAS_L.append(['Boadicea', 'Blue Sapphire Supercharged', 'S', 5, 7, 4, 7, 8, None, 'TEM', 'FRASCO',
        'Rosa-oud de saída cítrica-aldeídica e coração framboesa-cardamomo sobre oud-âmbar-patchouli: o rosa-oud mais luminoso da coleção.',
        'laranja, limão, aldeídos', 'rosa, cardamomo, jasmim (framboesa e ylang citados pela marca)', 'oud, âmbar, patchouli, toque de couro',
        'Rosa-oud luminoso | Noite Ameno–MF + Lazer Fresco–Frio | frasco ativo desde 09/out/2026; tier S dado por ele. Pirâmide: Selfridges, Neiman Marcus, descrição oficial Boadicea (verificada em 23/out/2026). Fragrantica: Oriental Floral, 2022. Registro franco: Rosal com 7 frascos (6 S) — redundância alta para a próxima rosa-oriental.',
        '', '', 'O Rosal', '—', '', 'Frasco Gui'])
if ('Mykonos', 'Empire') not in names:
    NOVAS_L.append(['Mykonos', 'Empire', 'A', 4, 7, 9, 7, 3, None, 'TEM', 'FRASCO',
        'Cítrico-tropical (lima, mandarina, água de coco) sobre patchouli-musk-tonka: frasco de calor e esporte.',
        'lima, bergamota, mandarina, frutas tropicais', 'laranja siciliana, gengibre, jasmim, água de coco', 'patchouli, musk branco, tonka',
        'Camada custo — só corrida/academia (slot E), fora de qualquer análise | tier A dado por ele (11/out/2026). Pirâmide de varejista (arabicparfums.com); Parfumo: híbrido de Afternoon Swim e Virgin Island Water.',
        '', '', 'O Marinho', '—', '', 'Frasco Gui'])
DA += NOVAS_L
DA.sort(key=lambda r: (r[0].lower(), r[1].lower()))
fg = sum(1 for r in DA if r[21] == 'Frasco Gui')
A = A[:k0] + json.dumps(DA, ensure_ascii=False) + A[k1:]
A = rep1(A, '<title>Avaliação de perfumes — v51 (tiers)</title>', '<title>Avaliação de perfumes — v52 (tiers)</title>')
h0 = A.index('<h1>Avaliação de perfumes — <em>'); h1 = A.index('</h1>', h0)
casas = len({r[0] for r in DA})
A = A[:h0] + f'<h1>Avaliação de perfumes — <em>{casas} casas</em> · {len(DA)} frascos · 23/out/2026 · v52 (tiers e posse espelhados da aba Tiers do Google Sheets após a limpeza de 23/out/2026 · entram Blue Sapphire Supercharged e Empire · Bedouin Rose C · {fg} frascos próprios conferidos)' + A[h1:]
open(f'out/avaliacao_perfumes_v52_2026-10-23.html', 'w', encoding='utf-8').write(A)
print('mapa v6 · observatório v4.8 · avaliação v52:', len(DA), 'linhas,', len(ch), 'tier/posse sincronizados,', len(NOVAS_L), 'novas · FG', fg, '· sem par:', len(miss))
print('  sem par na avaliação:', miss)

# ======== PARTE F — 00_LEIA_PRIMEIRO.md (cadeia v60) ========
L0 = open('00_LEIA_PRIMEIRO.md', encoding='utf-8').read()
txt = f"""# 00 — LEIA PRIMEIRO · fonte de verdade do projeto (23/out/2026 · cadeia v60)

> Vale para toda conversa deste projeto, inclusive quando a skill @perfumista for acionada. Este arquivo só traz **regras vigentes** — o histórico de cada decisão está nos próprios arquivos (seção *Histórico*) e não deve ser usado como regra.

## 1. Hierarquia de fontes

| Ordem | Fonte | Para quê |
|---|---|---|
| 1 | **Fala do Guilherme** na conversa | Vence qualquer arquivo. Posse de frasco se confirma com ele antes de presumir |
| 2 | **Google Sheets "Diário de Uso — Perfumes"** | Aba **Tiers** = fonte única de tier S/A/B/C/D (decisão dele, 04/out/2026). Aba **Diário** = registro de uso (slots M / N / E). Aba **Frascos** = último uso e contagens |
| 3 | **Arquivos do projeto (lista abaixo)** | Alocação, fichas, modelo diário, visualizações |
| 4 | Skill @perfumista (roteador) | Só o fluxo; aponta sempre para a versão mais alta de cada arquivo do projeto |

## 2. Arquivos do projeto — os únicos válidos

| Arquivo | O que é |
|---|---|
| `00_LEIA_PRIMEIRO.md` | Este arquivo |
| `playbook_v60.md` | Grade por faixa × ocasião: {len(cnt)} frascos, {TOTAL} alocações, Seguro / jogada / Também / sprays; custo por uso em dois slots |
| `fichas_tecnicas_v10.md` | Ficha de cada frasco: pirâmide, arquétipo, eixo, doçura, peso, Td, envelope, aptidão, sprays, janelas |
| `tabelao_v2_8.md` / `.xlsx` | Espelho da aba Tiers (355 linhas, posse + tier); xlsx com aba *Histórico* |
| `frascos_v8.py` + `modelo_diario_v1_7.py` | Modelo "Top 10 do dia" em dois slots (`python3 modelo_diario_v1_7.py --exemplo`) |
| `observatorio_colecao_v4_8.html` | Roda de arquétipos × tier, simulador por célula do playbook v60 |
| `mapa_climatico_colecao_v6.html` | Temperatura × Td × chuva, 71 frascos |
| `avaliacao_perfumes_v52_2026-10-23.html` | Base de avaliação (1.207 perfumes, pirâmides, veredictos históricos) |
| `gerador_v60.py` | Script que gera toda a cadeia v60 a partir da v59 + snapshot da aba Tiers, com verificações |

**Não existem mais e não devem ser citados:** qualquer versão anterior dos arquivos acima (playbook v59, fichas v9, tabelão v2.7, frascos_v7, modelo v1.6, observatório v4.6/v4.7, mapa v5, avaliação v51, gerador_v59), arquétipos v1.x, gaps v3.x, custo_por_uso_2slots, ranking_faixa, `claude/diario_uso.md`. O arquétipo de cada frasco está nas fichas; os gaps estão na seção 6; o custo por uso está no playbook.

## 3. Coleção

| Item | Valor |
|---|---|
| Grade | **{len(cnt)} frascos** — S {tiers_g['S']} · A {tiers_g['A']} · B {tiers_g['B']} · C {tiers_g['C']} (Oud Kyphi, Bleu de Chanel, Bedouin Rose) |
| Camada custo | 10 frascos — **usada só em academia e corrida; ignorada em qualquer análise ou framework** (regra dele, 25/set/2026). Aparece no modelo diário só com `Extra: corrida` ou `Extra: academia`. No diário, slot **E** |
| Total próprio | 81 frascos (= Frasco Gui na aba Tiers) |
| Frascos de 10 ml experimentais | Oud Kyphi, Bedouin Rose, K-Musk — não tratar como frasco cheio em custo por uso ou desbaste |

## 4. Regras vigentes

| Tema | Regra |
|---|---|
| **Classificação** | Só o tier dado por ele. Âncora: S compraria de novo amanhã · A alegria real · B respeito, cumpre função · C neutro · D fora. Nota numérica de perfume, "nota na pele", remedição, preliminar e contestado **não existem** (encerrado em 19/set/2026) |
| **Rank dentro do arquétipo** | Reflete o que ele mais usa e não precisa seguir o tier. Seguros e jogadas do playbook ficam por incumbência (decisão de 25/set/2026) |
| **Coleção** | **Reaberta desde 18/set/2026.** Nada de "PASSAR por default" |
| **Avaliação de compra** | 5 gates: DNA · gap · redundância · clima/custo por uso **por slot** · performance × preço. Postura dura; teste na pele antes de comprar (anti-blind-buy). **Doçura não reprova no gate de DNA** (fala dele, 18/out/2026: "não tenho mais essa demanda de reprovação de doce") — pesam gap, slot e redundância |
| **Rotina 2x/dia** | Manhã 7h00–7h30 (slot Dia, faixa pela máxima das horas restantes) e noite pós-banho 19h–20h (slot Noite, faixa pela temperatura de ~20–22h). Dois top 10 por dia. Não repetir o da manhã à noite; evitar dois densos (peso ≥4) no mesmo dia; em casa à noite −1 spray. À noite ele pode usar densos mesmo em casa — a intenção é curtir |
| **Longevidade** | O perfume do Dia **não precisa durar até a noite** (25/set/2026). Longevidade não é critério |
| **Custo por uso** | Noite Ameno/Fresco ≈ 300 noites/ano: noturno denso não reprova por clima, reprova por redundância no slot Noite e por tier. Frio/Muito Frio seguem escassos (~31 noites, ~5 dias). O gargalo é o slot Dia. Noite de calor é rara (~37/ano) |
| **Janelas decididas por ele** | **Vibrato:** só dia e trabalho (Lazer, Trab. Informal, Trab. Formal em MQ/Q/A), 9 janelas. **Tygar Extrait:** noites de MQ/Q/A. **Les Sables Roses:** noites exceto as quentes (Ameno–MF). **Bedouin Rose:** C na grade (23/out/2026) — só Cozy Noite no Frio/MF |
| **Janelas decididas pelo Claude a pedido dele** (23/out/2026, Fragrantica/Parfumo/marca — estimativa, não uso declarado) | **Favonius:** só noite Ameno–MF (N.Inf · N.For · Cozy N). **Omnia Omnibus Ubique:** só noite Ameno–MF (N.For e Cozy N 3, N.Inf 2). **Scarlet Sands:** Quente–Frio, Lazer + noites. **Blue Sapphire Supercharged:** noites Ameno–MF para sair + Lazer/Cozy N no Fresco/Frio. O diário confirma ou derruba |
| **Desbaste** | Os C ficam ("têm valor em circunstâncias"). B redundantes só se decide com ≥90 dias de diário 2x/dia |
| **Alteração da grade** | Só por script (gerador), nunca editando célula à mão. Reabre com: mudança de tier na aba Tiers, aptidão/envelope declarado por ele, frasco novo ou saída |
| **Descrição de notas** | Sempre com fonte (Fragrantica, Parfumo, site da marca); nunca inventar pirâmide |

## 5. Diário de uso

Vive na aba **Diário** do Google Sheets (primeiro registro: 19/set/2026). Coluna Data sempre em **texto** `AAAA-MM-DD`, em ordem cronológica (formatar a coluna A como texto simples — senão o Sheets converte em número serial). Colunas: Data · Slot · Perfume (nome igual ao da coluna "Nome no diário" da aba Tiers) · Ocasião · Temp · Td · Sprays · NotaDia · Obs.

| Campo | Regra |
|---|---|
| **Slot** | **M** = manhã (slot Dia) · **N** = noite pós-banho (slot Noite) · **E** = Esporte (camada custo, corrida/academia; decisão dele, 23/out/2026). Segunda aplicação no mesmo slot = nova linha com o mesmo slot e Obs "2ª aplicação" (não existe `N2`) |
| **NotaDia** | **Acerto do dia, 1–5** (decisão dele, 23/out/2026): o quanto aquele perfume serviu àquele dia/slot (clima, ocasião, humor). 5 acertou em cheio · 4 bom · 3 ok · 2 errado para o dia · 1 não repetir nessa condição. **Não é nota do perfume** — tier continua sendo a única régua de gosto. Uso: calibrar janelas do playbook (2–3 repetidos numa faixa/slot tiram o frasco da célula) e alimentar o desbaste de B redundantes. Não retroativa |

## 6. Gaps (status conforme observatório)

| # | Gap | Arquétipo | Status |
|---|---|---|---|
| 1 | Vetiver seco/mineral protagonista | Geólogo | aberto |
| 2 | Chypre verde aromático | Botânico | aberto |
| 3 | Tabaco dark | — | **fechado** (Montabaco Intensivo) |
| 4 | Sândalo cremoso protagonista (gap de dia/cozy) | Marceneiro | aberto — Santal 33 e Wood Whisper não fecham |
| 5 | Leather raw birch-tar | Curtidor | aberto, baixa prioridade |

**Território saturado (não é gap, é alerta):** rosa-oud / Rosal — 7 frascos na grade, 6 S. Compra de rosa-oriental parte de redundância alta no gate 3.

## 7. Pendências abertas

| # | Pendência | Com quem |
|---|---|---|
| 1 | Validar no diário (NotaDia) as janelas de Favonius, Omnia Omnibus Ubique, Scarlet Sands e Blue Sapphire Supercharged — decididas pelo Claude em 23/out/2026 | Guilherme, pelo uso |
| 2 | Divisão do arquétipo Mediterrâneo em Mediterrâneo + Solar (Imagination, Lucius, Tygar Extrait), pedida em 05/out/2026: não propagada a fichas, mapa e observatório | Guilherme decide se segue |
| 3 | Amouage Enclave sem tier (ainda vai avaliar) | Guilherme |
"""
open('out/00_LEIA_PRIMEIRO.md', 'w', encoding='utf-8').write(txt)
print('LEIA_PRIMEIRO v60 escrito')
